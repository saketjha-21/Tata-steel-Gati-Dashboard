from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR

# Brand-inspired palette (blue-white dominant)
NAVY = RGBColor(0, 51, 102)
BLUE = RGBColor(0, 102, 179)
LIGHT_BLUE = RGBColor(229, 240, 250)
ACCENT = RGBColor(0, 153, 204)
GREY = RGBColor(90, 90, 90)
LIGHT_GREY = RGBColor(245, 247, 250)
WHITE = RGBColor(255, 255, 255)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def add_header(slide, title):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.65))
    bar.fill.solid(); bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    tx = bar.text_frame
    tx.text = title
    p = tx.paragraphs[0]
    p.font.size = Pt(20); p.font.bold = True; p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.LEFT


def add_textbox(slide, x,y,w,h,text,size=14,bold=False,color=GREY):
    box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=box.text_frame; tf.clear()
    p=tf.paragraphs[0]; p.text=text
    p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=color
    return box

# Title slide
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_header(slide, "Observations & Organizational Analysis")
add_textbox(slide,0.8,1.4,8,1,"Management Trainee Review",24,True,NAVY)
add_textbox(slide,0.8,2.2,8,1.2,"Krsnaa Diagnostics Ltd\nSaket Jha",16,False,GREY)
icon = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.1), Inches(1.5), Inches(3.6), Inches(4.8))
icon.fill.solid(); icon.fill.fore_color.rgb = LIGHT_BLUE; icon.line.color.rgb = BLUE
icon.text_frame.text = "Healthcare +\nOperations\nOverview"
for p in icon.text_frame.paragraphs:
    p.font.size = Pt(18); p.font.bold=True; p.alignment=PP_ALIGN.CENTER; p.font.color.rgb=NAVY

# Slide1
slide=prs.slides.add_slide(prs.slide_layouts[6]); add_header(slide,"1. Underlying Organizational Frictions")
box=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.6), Inches(0.9), Inches(4.5), Inches(1.6))
box.fill.solid(); box.fill.fore_color.rgb=LIGHT_GREY; box.line.color.rgb=BLUE
box.text_frame.text="Factors largely outside operational influence\n• Long receivable cycles\n• Dependence on PPP model\n• Limited retail contribution\n• PAT growth inconsistency"
for i,p in enumerate(box.text_frame.paragraphs):
    p.font.size=Pt(11 if i else 12); p.font.bold=(i==0); p.font.color.rgb=GREY
add_textbox(slide,0.8,1.0,7.5,0.6,"Observed challenges appear linked to fragmented execution structures",16,True,NAVY)
labels=["Technology\nfragmentation","Multiple\nsystems","Manual\nworkarounds","Process\ninconsistency","Operational\nvariability"]
x=0.9
for lab in labels:
    sh=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.3), Inches(2.2), Inches(1.0))
    sh.fill.solid(); sh.fill.fore_color.rgb=WHITE; sh.line.color.rgb=BLUE
    sh.text_frame.text=lab
    for p in sh.text_frame.paragraphs: p.alignment=PP_ALIGN.CENTER; p.font.size=Pt(12); p.font.bold=True; p.font.color.rgb=NAVY
    x+=2.45
for i in range(4):
    ln=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(3.1+i*2.45), Inches(2.8), Inches(3.35+i*2.45), Inches(2.8))
    ln.line.color.rgb=ACCENT; ln.line.width=Pt(2)
add_textbox(slide,0.9,3.7,12,0.8,"• Slow or incomplete automation implementation    • Teams relying on multiple systems    • Gaps between systems and workflows    • Software adoption inconsistency",12)
ins=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.4), Inches(12), Inches(0.6))
ins.fill.solid(); ins.fill.fore_color.rgb=LIGHT_BLUE; ins.line.fill.background()
ins.text_frame.text="Recurring theme across departments: technology exists, but implementation depth varies."
ins.text_frame.paragraphs[0].font.size=Pt(12); ins.text_frame.paragraphs[0].font.bold=True; ins.text_frame.paragraphs[0].font.color.rgb=NAVY

# Slide2
slide=prs.slides.add_slide(prs.slide_layouts[6]); add_header(slide,"2. Execution Inefficiencies + Resistance to Change")
add_textbox(slide,0.8,1.0,6,0.5,"Operational Inefficiencies",16,True,NAVY)
steps=["Duplicate efforts","Routine task escalation","Decision layers","Uneven ownership"]
for i,s in enumerate(steps):
    y=1.6+i*1.1
    sh=slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(0.8), Inches(y), Inches(5.8), Inches(0.8))
    sh.fill.solid(); sh.fill.fore_color.rgb=LIGHT_BLUE; sh.line.color.rgb=BLUE
    sh.text_frame.text=s
    p=sh.text_frame.paragraphs[0]; p.font.size=Pt(12); p.font.bold=True; p.font.color.rgb=NAVY
add_textbox(slide,6.9,1.2,5.5,0.5,"Resistance to Change",16,True,NAVY)
tri=slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, Inches(8.0), Inches(2.0), Inches(3.5), Inches(3.0))
tri.fill.solid(); tri.fill.fore_color.rgb=LIGHT_GREY; tri.line.color.rgb=BLUE
add_textbox(slide,9.2,2.2,2,0.4,"Leadership",11,True)
add_textbox(slide,8.7,3.0,3,0.4,"Middle management",11,True)
add_textbox(slide,9.0,3.8,2.5,0.4,"Ground teams",11,True)
add_textbox(slide,6.9,5.2,5.8,1.2,"• Information silos\n• Change uncertainty\n• Limited cross-level alignment\n• Need stronger communication loops",11)
add_textbox(slide,7.0,4.7,5.5,0.4,'"Different perspectives create execution gaps"',10,False,ACCENT)
add_textbox(slide,0.8,6.7,12,0.4,"Execution barriers often appear behavioral as much as structural.",11,True,GREY)

# Slide3
slide=prs.slides.add_slide(prs.slide_layouts[6]); add_header(slide,"3. Revenue Variability and Referral Dependence")
add_textbox(slide,0.8,1.0,12,0.5,"Infrastructure strength is not translating uniformly into utilization",16,True,NAVY)
funnel=[("Machine utilization",4.5), ("Referral flow",3.8), ("Patient inflow",3.1), ("Revenue consistency",2.4)]
y=1.8
for t,w in funnel:
    sh=slide.shapes.add_shape(MSO_SHAPE.TRAPEZOID, Inches(0.9+(4.5-w)/2), Inches(y), Inches(w), Inches(0.75))
    sh.fill.solid(); sh.fill.fore_color.rgb=LIGHT_BLUE; sh.line.color.rgb=BLUE
    sh.text_frame.text=t
    p=sh.text_frame.paragraphs[0]; p.font.size=Pt(12); p.font.bold=True; p.alignment=PP_ALIGN.CENTER; p.font.color.rgb=NAVY
    y+=0.8
panel=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.0), Inches(1.8), Inches(6.5), Inches(4.6))
panel.fill.solid(); panel.fill.fore_color.rgb=WHITE; panel.line.color.rgb=BLUE
panel.text_frame.text="Doctor Referral Challenge\n\nCurrent market perception often positions Krsnaa primarily as affordability-led.\n\nDirectional ideas:\n• Clinical partnerships\n• Doctor ecosystem engagement\n• Digital consultation alliances\n• Sales capability enhancement"
for i,p in enumerate(panel.text_frame.paragraphs): p.font.size=Pt(12 if i==0 else 11); p.font.bold=(i==0); p.font.color.rgb=(NAVY if i==0 else GREY)
add_textbox(slide,0.8,6.6,12,0.4,"Current ideas are directional observations rather than recommendations.",11,True,GREY)

# Slide4
slide=prs.slides.add_slide(prs.slide_layouts[6]); add_header(slide,"4. Retail Presence and Market Awareness")
add_textbox(slide,0.8,1.0,12,0.5,"Infrastructure scale exists ahead of brand recall",16,True,NAVY)
left=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.1), Inches(3.8))
left.fill.solid(); left.fill.fore_color.rgb=LIGHT_GREY; left.line.color.rgb=BLUE
left.text_frame.text="Retail challenge\n• Limited public awareness\n• Strong local players\n• Dependence on partnerships"
right=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.1), Inches(1.8), Inches(5.4), Inches(3.8))
right.fill.solid(); right.fill.fore_color.rgb=WHITE; right.line.color.rgb=BLUE
right.text_frame.text="Marketing observations\n• Awareness efforts still evolving\n• Opportunity in radiology positioning\n• Local market visibility opportunities\n• External support may accelerate execution"
for shp in [left,right]:
    for i,p in enumerate(shp.text_frame.paragraphs): p.font.size=Pt(12 if i==0 else 11); p.font.bold=(i==0); p.font.color.rgb=(NAVY if i==0 else GREY)
ln=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(5.9), Inches(3.6), Inches(7.0), Inches(3.6)); ln.line.color.rgb=ACCENT; ln.line.width=Pt(2.5)
add_textbox(slide,0.8,6.5,12,0.5,"Growth infrastructure exists; awareness and positioning remain the constraint.",12,True,GREY)

# Slide5
slide=prs.slides.add_slide(prs.slide_layouts[6]); add_header(slide,"5. Scorecard Analysis — Key Findings")
metric=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.2), Inches(4.0), Inches(2.2))
metric.fill.solid(); metric.fill.fore_color.rgb=NAVY; metric.line.fill.background()
metric.text_frame.text="Revenue Risk\nContribution\n42%"
for i,p in enumerate(metric.text_frame.paragraphs): p.font.size=Pt(16 if i<2 else 30); p.font.bold=True; p.font.color.rgb=WHITE; p.alignment=PP_ALIGN.CENTER
add_textbox(slide,0.9,3.7,3.8,0.5,"Material Cost: 11.6%",13,True,NAVY)
add_textbox(slide,0.9,4.2,3.8,0.5,"Process KPIs: 7.1%",13,True,NAVY)
chart=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(5.2), Inches(1.2), Inches(3.1), Inches(3.6))
chart.fill.solid(); chart.fill.fore_color.rgb=LIGHT_GREY; chart.line.color.rgb=BLUE
chart.text_frame.text="Variability by role\n\nDistrict Coordinator CV: 0.58\nCluster Manager CV: 0.37"
for i,p in enumerate(chart.text_frame.paragraphs): p.font.size=Pt(12 if i==0 else 11); p.font.bold=(i==0); p.font.color.rgb=GREY
find=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.6), Inches(1.2), Inches(4.0), Inches(4.9))
find.fill.solid(); find.fill.fore_color.rgb=WHITE; find.line.color.rgb=BLUE
find.text_frame.text="Findings\n• High variability rather than weak capability\n• Revenue performance differs significantly across centers\n• Technology adoption instability observed\n• Execution often dependent on individuals"
for i,p in enumerate(find.text_frame.paragraphs): p.font.size=Pt(12 if i==0 else 11); p.font.bold=(i==0); p.font.color.rgb=(NAVY if i==0 else GREY)
add_textbox(slide,0.8,6.5,12,0.4,"Challenge appears to be standardization and consistency—not capability.",12,True,GREY)

# Slide6
slide=prs.slides.add_slide(prs.slide_layouts[6]); add_header(slide,"6. Summary and Management Trainee Observations")
for i,t in enumerate(["Technology","Execution","Revenue"]):
    sh=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2+i*3.4), Inches(1.5), Inches(2.7), Inches(1.0))
    sh.fill.solid(); sh.fill.fore_color.rgb=LIGHT_BLUE; sh.line.color.rgb=BLUE
    sh.text_frame.text=t
    p=sh.text_frame.paragraphs[0]; p.font.bold=True; p.font.size=Pt(14); p.alignment=PP_ALIGN.CENTER; p.font.color.rgb=NAVY
cons=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.5), Inches(1.4), Inches(2.6), Inches(1.2))
cons.fill.solid(); cons.fill.fore_color.rgb=NAVY; cons.line.fill.background(); cons.text_frame.text="Consistency"
cp=cons.text_frame.paragraphs[0]; cp.font.size=Pt(18); cp.font.bold=True; cp.font.color.rgb=WHITE; cp.alignment=PP_ALIGN.CENTER
for x in [3.9,7.3,10.7]:
    ln=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(2.0), Inches(9.5), Inches(2.0)); ln.line.color.rgb=ACCENT; ln.line.width=Pt(2)
obs=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(3.1), Inches(12.3), Inches(2.8))
obs.fill.solid(); obs.fill.fore_color.rgb=WHITE; obs.line.color.rgb=BLUE
obs.text_frame.text="Management trainee observations\n• Program structure remained highly flexible\n• Practical exposure varied by department\n• Department expectations were not always aligned\n• Timelines and ownership structures could improve"
for i,p in enumerate(obs.text_frame.paragraphs): p.font.size=Pt(12 if i==0 else 11); p.font.bold=(i==0); p.font.color.rgb=(NAVY if i==0 else GREY)
add_textbox(slide,0.8,6.4,12,0.5,"The organization appears fundamentally strong; greater standardization may unlock more consistent execution at scale.",12,True,GREY)

out='Krsnaa_Leadership_Presentation.pptx'
prs.save(out)
print(f'Created {out}')
